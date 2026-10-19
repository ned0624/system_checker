# -*- coding: utf-8 -*-
"""
可忽略事件ID 搜尋小工具（CLI，獨立程式，不需要 config.json）

1. 執行後輸入可忽略清單的 Excel 路徑（可直接把檔案拖進視窗；直接按 Enter 沿用上次的檔案）
2. 輸入 ID 查詢，一次可輸入多筆，用逗號隔開（, ，、 都可以），
   逐筆顯示是否在可忽略清單，在清單內的會一併列出該列的其他欄位（例如說明）

指令:
    r / reload        重新讀取 Excel（清單有更新時用）
    f / file          換一個 Excel 檔
    q / quit / exit   結束程式（也可以直接關閉視窗）

清單位置: 工作表「可忽略事件ID」（找不到時列出工作表讓你選），
          ID 欄依標題「可忽略事件ID」或「ID」尋找，都沒有時用第一欄

需求套件: pip install openpyxl（不用 pandas，打包後的 exe 小、啟動快）
打包: 執行 build_exe.bat（排除用不到的 pandas/numpy/PIL/lxml，exe 小、啟動快）
"""
import os
import re
import sys
import openpyxl

# 輸出被導向檔案或非中文系統時，顯示不了的字以 ? 代替，不讓程式中斷
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(errors="replace")
    except (AttributeError, ValueError):
        pass

DEFAULT_SHEET = "可忽略事件ID"
ID_HEADERS = ("可忽略事件ID", "ID", "Id", "事件識別碼")
QUIT_CMDS = {"q", "quit", "exit", "離開"}
RELOAD_CMDS = {"r", "reload"}
FILE_CMDS = {"f", "file"}

def _base_dir():
    """程式所在資料夾；包成 exe 時是 exe 所在資料夾"""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


# 記住上次的 Excel 路徑（放在程式旁邊）
LAST_PATH_FILE = os.path.join(_base_dir(), "search_ignore_last.txt")


class Quit(Exception):
    pass


def ask(prompt):
    try:
        text = input(prompt).replace("\ufeff", "").strip()  # 去掉管線輸入可能夾帶的 BOM
    except (EOFError, KeyboardInterrupt):
        print()
        raise Quit
    if text.lower() in QUIT_CMDS:
        raise Quit
    return text


def is_blank(v):
    return v is None or str(v).strip() == ""


def norm(v):
    """統一 ID 格式：去空白、去掉 Excel 讀成 float 的 .0"""
    if is_blank(v):
        return None
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s or None


def read_last_path():
    try:
        with open(LAST_PATH_FILE, encoding="utf-8") as f:
            return f.read().strip()
    except OSError:
        return ""


def save_last_path(path):
    try:
        with open(LAST_PATH_FILE, "w", encoding="utf-8") as f:
            f.write(path)
    except OSError:
        pass


def ask_excel():
    """詢問 Excel 路徑，直到找到檔案為止"""
    last = read_last_path()
    while True:
        hint = f" (Enter = {last})" if last else ""
        path = ask(f"\nIgnore list Excel path{hint}> ").strip('"').strip("'")
        if not path and last:
            path = last
        if not path:
            continue
        if os.path.isfile(path):
            return path
        print(f"[ERROR] File not found: {path}")


def pick_sheet(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    sheets = wb.sheetnames
    wb.close()
    if DEFAULT_SHEET in sheets:
        return DEFAULT_SHEET
    print(f"Sheet '{DEFAULT_SHEET}' not found. Choose a sheet:")
    for i, name in enumerate(sheets, 1):
        print(f"  {i}. {name}")
    while True:
        text = ask("Sheet number> ")
        if text.isdigit() and 1 <= int(text) <= len(sheets):
            return sheets[int(text) - 1]


def load_ignore_table(path, sheet):
    """回傳 ID → 該列其他欄位的說明文字（一格多個 ID 時每個 ID 都對應同一列）"""
    # read_only + data_only：只讀值（公式取計算結果），大檔也快
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb[sheet].iter_rows(values_only=True)
        header = next(rows, ())
        headers = [None if is_blank(h) else str(h).strip() for h in header]
        names = [h for h in headers if h]
        id_name = next((h for h in ID_HEADERS if h in names), None)
        id_idx = headers.index(id_name) if id_name else 0
        id_col = id_name or "first column"
        others = [(i, h) for i, h in enumerate(headers) if h and i != id_idx]
        table = {}
        for rownum, row in enumerate(rows, start=2):
            v = row[id_idx] if id_idx < len(row) else None
            if is_blank(v):
                continue
            info = "  ".join(f"{h}: {str(row[i]).strip()}" for i, h in others
                             if i < len(row) and not is_blank(row[i]))
            info = f"row {rownum}" + (f"  {info}" if info else "")
            for part in str(v).replace("\n", ",").replace("、", ",").split(","):
                n = norm(part)
                if n:
                    table.setdefault(n, []).append(info)
    finally:
        wb.close()
    print(f"Loaded {len(table)} ignorable event IDs ({os.path.basename(path)} / {sheet} / column: {id_col})")
    return table


def split_ids(text):
    seen, ids = set(), []
    for part in re.split(r"[,，、;\s]+", text):
        n = norm(part)
        if n and n not in seen:
            seen.add(n)
            ids.append(n)
    return ids


def search(table, text):
    ids = split_ids(text)
    if not ids:
        return
    hit = 0
    for eid in ids:
        if eid in table:
            hit += 1
            print(f"  [IGNORABLE] {eid}")
            for info in table[eid]:
                print(f"                {info}")
        else:
            print(f"  [NOT IN LIST] {eid}")
    print(f"Total {len(ids)}: ignorable {hit}, not in list {len(ids) - hit}")


def open_excel():
    """選檔並載入，失敗就重新選，回傳 (path, sheet, table)"""
    while True:
        path = ask_excel()
        try:
            sheet = pick_sheet(path)
            table = load_ignore_table(path, sheet)
        except Quit:
            raise
        except Exception as e:
            print(f"[ERROR] Cannot read file (must be .xlsx; it can stay open in Excel): {e}")
            continue
        save_last_path(path)
        return path, sheet, table


def main():
    print("Ignorable Event ID Search")
    print("Enter IDs separated by commas | r = reload | f = change file | q = quit (or close the window)")
    try:
        path, sheet, table = open_excel()
        while True:
            text = ask("\nID> ")
            cmd = text.lower()
            if cmd in RELOAD_CMDS:
                try:
                    table = load_ignore_table(path, sheet)
                except Exception as e:
                    print(f"[ERROR] Reload failed, keeping previous list: {e}")
            elif cmd in FILE_CMDS:
                path, sheet, table = open_excel()
            else:
                search(table, text)
    except Quit:
        pass


if __name__ == "__main__":
    main()
