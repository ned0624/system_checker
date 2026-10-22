# -*- coding: utf-8 -*-
"""
週檢小幫手：依 CSV 檢核結果填寫系統檢核表，並把 CSV 內容貼到問題追蹤表

流程:
1. 讀取「系統檢核問題追蹤.xlsx」的「可忽略事件ID」清單
2. 讀取指定資料夾內的 CSV，檔名 = [日期_]伺服器名（例: 20261015_P3-WFA-AP01V.csv）
   資料夾內有多個日期時，只處理日期最新的那一批
   只讀「錯誤」「警告」（以及「重大」，算錯誤）的事件，「資訊」等其他等級全部略過
3. 檢核表（BOS_SystemCheck.xlsx）：到第一欄找到同名伺服器那一列
   - CSV 裡的 ID 全部都在可忽略清單 → 「警告」「錯誤」兩欄都寫 O
   - 有 ID 不在清單 → 依 CSV 的等級欄位（警告/錯誤）在對應欄打 X，
     另一欄若沒有問題則寫 O；不在清單的事件說明寫進「備註」
4. 問題追蹤表（系統檢核問題追蹤.xlsx）：CSV 的錯誤/警告事件（含可忽略的事件）接在最後一筆資料下面
   依標題找欄位（工作表要先有標題列）：
   - Server Name：檔名中的伺服器名
   - 周：依檔名日期算 ISO 週次，格式 W+年後兩碼+週次（20261015 → W2642）；檔名沒有日期時用今天
   - 錯誤/警告、事件識別碼、記錄檔名稱、事件內容：依 track.columns 對應 CSV 欄位
     （預設 LevelDisplayName / Id / LogName / Message）
   - 可忽略：填入 COUNTIF 公式，事件識別碼在「可忽略事件ID」清單就顯示 V
   - 可能造成的影響、Action 等其他欄不動
   - 同一台伺服器、同一週別的舊資料會先刪除再貼上，重跑不會重複

需求套件: pip install pandas openpyxl
範本模式（config.json 的 check_sheet.template 有設定時）:
   - 每週產生完整檢核表到 check_sheet.output（路徑可用 {week}=W2642、{yy}=26、{ww}=42，資料夾會自動建立）
   - 當週檔已存在（重跑）就以它為底更新，保留手動填的內容；否則從範本產生
   - 有底色的格子（灰、黑、黃…）與合併儲存格不填
   - 沒有本週 CSV 的伺服器，事件檢視器兩欄填 O O；防毒軟體欄全填 O
   - 連結欄：有設定 check_sheet.link_target 時，每一列的超連結由程式產生
     （例: //server/share/System Check/{week}/Reference/BOS，可用 {week}/{yy}/{ww}/{host}）；
     link_text 設定顯示文字，null = 沿用範本原本文字並換週別（空白格顯示最後一層資料夾名稱）
     沒設定 link_target 時，只把範本既有連結的 W2640 → 本週 W2642、_W40 → _W42
   - 只處理本週（或 --week 指定週）的 CSV，同一台有多次取最新日期

用法:
    python fill_check.py [--week W2642]
    python fill_check.py <csv資料夾> <檢核表xlsx> <追蹤表xlsx>   （有給的參數會取代設定檔的路徑）

設定檔: 與本程式同資料夾的 config.json
    paths.csv_dir                             要讀取的 CSV 資料夾
    paths.check_xlsx                          要填寫 O/X 的檢核表（BOS_SystemCheck.xlsx）
                                              路徑中可寫 {week}，會換成 CSV 的週別（例: BOS_..._{week}.xlsx → _W2642.xlsx）
    paths.track_xlsx                          問題追蹤表（系統檢核問題追蹤.xlsx），可忽略清單也從這裡讀
    ignore_list.sheet / id_column             可忽略事件ID 的工作表與欄位
    csv.id_column / level_column / desc_column  CSV 的 ID、警告/錯誤分類、說明欄位
    check_sheet.sheet                         工作表名稱，null = 第一個工作表
    check_sheet.header_row                    標題列在第幾列
    check_sheet.warn/error/note_header        警告、錯誤、備註欄的標題
    check_sheet.output                        另存路徑，null = 直接覆寫檢核表
    track.sheet                               工作表名稱；null = 第一個工作表
    track.header_row                          標題列在第幾列
    track.host_header / week_header           伺服器、週別欄的標題
    track.columns                             追蹤表標題 → CSV 欄位名稱
    track.ignore_header                       可忽略欄的標題，null = 不填
    track.output                              另存路徑，null = 直接覆寫追蹤表
    log_dir                                   log 資料夾，省略 = logs（依日期命名，例: 20261016.log；同一天多次執行附加在同一檔）
    路徑可用絕對路徑（例: D:/週檢/BOS_SystemCheck.xlsx 或 \\\\server\\share\\週檢\\csv），
    JSON 裡的反斜線要寫兩次（\\），或改用正斜線 /
    相對路徑以 config.json 所在資料夾為基準
"""
import sys
import glob
import json
import os
import re
import traceback
from datetime import date, datetime
import pandas as pd
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font

# ===== 讀取設定檔 =====
# 輸出被導向檔案或非中文系統時，顯示不了的字以 ? 代替，不讓程式中斷
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(errors="replace")
    except (AttributeError, ValueError):
        pass


def _base_dir():
    """程式所在資料夾；包成執行檔（PyInstaller）時是 exe 所在資料夾，config.json 放在旁邊"""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


BASE_DIR = _base_dir()
CONFIG_PATH = os.path.join(BASE_DIR, "config.json")


def _path(p):
    return os.path.join(BASE_DIR, p) if p else p


try:
    with open(CONFIG_PATH, encoding="utf-8-sig") as f:
        _cfg = json.load(f)
except Exception as _e:
    if __name__ != "__main__":
        raise
    # 設定檔讀不到時還不知道 log_dir，log 寫在預設的 logs 資料夾
    _msg = f"[ERROR] Cannot read config file {CONFIG_PATH}: {type(_e).__name__}: {_e}"
    print(_msg)
    os.makedirs(os.path.join(BASE_DIR, "logs"), exist_ok=True)
    with open(os.path.join(BASE_DIR, "logs", f"{datetime.now():%Y%m%d}.log"),
              "a", encoding="utf-8") as _f:
        _f.write(f"\n===== {datetime.now():%Y-%m-%d %H:%M:%S} =====\n" + _msg + "\n\n" + traceback.format_exc())
    sys.exit(1)

# 舊版設定檔沒有 paths 區段時，沿用 csv.dir / check_sheet.xlsx / track.xlsx
_paths = _cfg.get("paths", {})
CSV_DIR    = _path(_paths.get("csv_dir") or _cfg["csv"].get("dir"))
CHECK_XLSX = _path(_paths.get("check_xlsx") or _cfg["check_sheet"].get("xlsx"))
TRACK_XLSX = _path(_paths.get("track_xlsx") or _cfg["track"].get("xlsx"))

IGNORE_XLSX   = TRACK_XLSX
IGNORE_SHEET  = _cfg["ignore_list"]["sheet"]
IGNORE_COLUMN = _cfg["ignore_list"]["id_column"]

CSV_ID_COLUMN    = _cfg["csv"]["id_column"]
CSV_LEVEL_COLUMN = _cfg["csv"]["level_column"]
CSV_DESC_COLUMN  = _cfg["csv"]["desc_column"]

CHECK_SHEET  = _cfg["check_sheet"]["sheet"]
HEADER_ROW   = _cfg["check_sheet"]["header_row"]
WARN_HEADER  = _cfg["check_sheet"]["warn_header"]
ERROR_HEADER = _cfg["check_sheet"]["error_header"]
NOTE_HEADER  = _cfg["check_sheet"]["note_header"]
OUTPUT_XLSX  = _path(_cfg["check_sheet"]["output"])
# 範本模式：有設定 template 時，每週從範本（或已存在的當週檔）產生完整檢核表存到 output
TEMPLATE_XLSX    = _path(_cfg["check_sheet"].get("template"))
ANTIVIRUS_HEADER = _cfg["check_sheet"].get("antivirus_header")
LINK_HEADER      = _cfg["check_sheet"].get("link_header")
# 連結規則：有設定 link_target 時，每一列的超連結由程式產生（可用 {week}/{yy}/{ww}/{host}）
LINK_TARGET      = _cfg["check_sheet"].get("link_target")
LINK_TEXT        = _cfg["check_sheet"].get("link_text")

TRACK_SHEET      = _cfg["track"]["sheet"]
TRACK_HEADER_ROW = _cfg["track"].get("header_row", 1)
HOST_HEADER      = _cfg["track"]["host_header"]
WEEK_HEADER      = _cfg["track"]["week_header"]
TRACK_COLUMNS    = _cfg["track"].get("columns") or {
    "錯誤/警告": CSV_LEVEL_COLUMN, "事件識別碼": CSV_ID_COLUMN,
    "記錄檔名稱": "LogName", "事件內容": CSV_DESC_COLUMN}
IGNORE_HEADER    = _cfg["track"].get("ignore_header", "可忽略")
# 追蹤表中放事件 ID 的欄（可忽略公式用）
TRACK_ID_HEADER  = next((h for h, c in TRACK_COLUMNS.items() if c == CSV_ID_COLUMN), None)
TRACK_OUTPUT = _path(_cfg["track"]["output"])

LOG_DIR = _path(_cfg.get("log_dir") or "logs")
# ======================

# 事件等級（CSV 等級欄位的值，不分大小寫）；不在這兩組的等級（例: 資訊）一律不讀
ERROR_LEVELS = {"錯誤", "重大", "error", "critical"}
WARN_LEVELS  = {"警告", "warning"}
# 確定可以略過的等級；其他沒見過的值（例如編碼錯誤的亂碼）會提出警告，避免錯誤被默默略過
SKIP_LEVELS  = {"資訊", "詳細資訊", "詳細", "information", "verbose", ""}


def level_kind(level):
    """回傳 "error" / "warn" / None（資訊等其他等級）"""
    lv = str(level).strip().lower() if pd.notna(level) else ""
    if lv in ERROR_LEVELS:
        return "error"
    if lv in WARN_LEVELS:
        return "warn"
    return None


def keep_warn_error(df):
    """只留下錯誤/警告的列，回傳 (過濾後的 df, 略過筆數, 無法辨識的等級值)"""
    if CSV_LEVEL_COLUMN not in df.columns:
        raise KeyError(f"Column '{CSV_LEVEL_COLUMN}' not found, columns: {list(df.columns)}")
    levels = df[CSV_LEVEL_COLUMN]
    mask = levels.map(level_kind).notna()
    unknown = sorted({str(v).strip() for v in levels[~mask].dropna()
                      if str(v).strip().lower() not in SKIP_LEVELS})
    return df[mask].reset_index(drop=True), int((~mask).sum()), unknown


def norm(v):
    """統一 ID 格式：去空白、去掉 Excel 讀成 float 的 .0"""
    if pd.isna(v):
        return None
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s or None


def read_csv_any(path):
    """依序嘗試常見編碼（含 Big5/CP950）"""
    for enc in ("utf-8-sig", "cp950", "big5", "utf-8"):
        try:
            df = pd.read_csv(path, dtype=str, encoding=enc)
            df.columns = df.columns.str.strip()
            return df
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Cannot detect file encoding: {path}")


def load_ignore_ids(xlsx_path):
    df = pd.read_excel(xlsx_path, sheet_name=IGNORE_SHEET, dtype=str)
    df.columns = df.columns.str.strip()
    if IGNORE_COLUMN not in df.columns:
        raise KeyError(f"Ignore list: column '{IGNORE_COLUMN}' not found, columns: {list(df.columns)}")
    ids = set()
    for v in df[IGNORE_COLUMN]:
        # 若一格內有多個 ID（以逗號/換行分隔）也一併拆開
        if pd.isna(v):
            continue
        for part in str(v).replace("\n", ",").replace("、", ",").split(","):
            n = norm(part)
            if n:
                ids.add(n)
    return ids


# ---------- 檢核表 O/X ----------

def find_columns(ws, optional=()):
    """依標題文字找出 警告/錯誤/備註（以及 optional 中有設定的欄）的欄號"""
    headers = {str(c.value).strip(): c.column for c in ws[HEADER_ROW] if c.value is not None}
    cols = {}
    for h in (WARN_HEADER, ERROR_HEADER, NOTE_HEADER) + tuple(h for h in optional if h):
        if h not in headers:
            raise KeyError(f"Check sheet: header '{h}' not found, headers: {list(headers)}")
        cols[h] = headers[h]
    return cols


def is_shaded(cell):
    """有底色（灰、黑、黃…）或是合併儲存格 → True，這種格子不填"""
    if isinstance(cell, MergedCell):
        return True
    fill = cell.fill
    if fill is None or fill.fill_type in (None, "none"):
        return False
    if fill.fill_type != "solid":
        return True  # 花紋底色
    c = fill.fgColor
    if c is None:
        return False
    if c.type == "rgb":
        return str(c.rgb).upper()[-6:] != "FFFFFF"
    if c.type == "theme":
        return not (c.theme == 0 and not c.tint)  # theme 0 = 白色
    if c.type == "indexed":
        return c.indexed not in (9, 64)  # 9 = 白色, 64 = 系統預設
    return True


def set_cell(ws, row, col, value):
    """格子沒有底色才寫入，回傳是否有寫"""
    cell = ws.cell(row=row, column=col)
    if is_shaded(cell):
        return False
    cell.value = value
    return True


def update_week_tokens(text, week):
    """W2640 → W2642、_W40 → _W42（week 例: W2642）"""
    text = re.sub(r"W\d{4}(?!\d)", week, text)
    return re.sub(r"(?<!\d)W\d{2}(?!\d)", "W" + week[3:5], text)


def week_path(path, week):
    """路徑中的 {week}/{yy}/{ww} 換成週別，例: W2642 / 26 / 42"""
    return path.replace("{week}", week).replace("{yy}", week[1:3]).replace("{ww}", week[3:5])


def find_host_rows(ws):
    """第一欄的伺服器名稱 → 列號（不分大小寫）"""
    rows = {}
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is not None and str(v).strip():
            rows[str(v).strip().upper()] = r
    return rows


def check_csv(df, ignore_ids):
    """回傳 (有警告, 有錯誤, 備註文字)；df 已只含錯誤/警告，只計算不在可忽略清單的事件"""
    for col in (CSV_ID_COLUMN, CSV_LEVEL_COLUMN, CSV_DESC_COLUMN):
        if col not in df.columns:
            raise KeyError(f"Column '{col}' not found, columns: {list(df.columns)}")

    has_warn = has_error = False
    notes = []
    seen = set()
    for _, row in df.iterrows():
        eid = norm(row[CSV_ID_COLUMN])
        if not eid or eid in ignore_ids:
            continue
        level = str(row[CSV_LEVEL_COLUMN]).strip() if pd.notna(row[CSV_LEVEL_COLUMN]) else ""
        desc = str(row[CSV_DESC_COLUMN]).strip() if pd.notna(row[CSV_DESC_COLUMN]) else ""
        if level_kind(level) == "error":
            has_error = True
        else:
            has_warn = True
        key = (level, eid, desc)
        if key not in seen:  # 同一事件重複出現只記一次
            seen.add(key)
            notes.append(f"[{level}] {eid} {desc}".strip())
    return has_warn, has_error, "\n".join(notes)


def fill_check_row(ws, cols, r, df, ignore_ids, host):
    has_warn, has_error, note = check_csv(df, ignore_ids)
    set_cell(ws, r, cols[WARN_HEADER], "X" if has_warn else "O")
    set_cell(ws, r, cols[ERROR_HEADER], "X" if has_error else "O")
    if set_cell(ws, r, cols[NOTE_HEADER], note or None) and note:
        ws.cell(row=r, column=cols[NOTE_HEADER]).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    print(f"  Check sheet: warning={'X' if has_warn else 'O'} error={'X' if has_error else 'O'}"
          + (f"\n    {note.replace(chr(10), chr(10) + '    ')}" if note else ""))


HOSTNAME_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]*$")
HYPERLINK_FORMULA_TEXT = re.compile(r',\s*"([^"]*)"\s*\)\s*$')


def set_link(cell, host, week):
    """依 link_target / link_text 規則設定超連結；link_text 為 null 時沿用格子原本文字（換週別）"""
    render = lambda s: week_path(s, week).replace("{host}", host)
    target = render(LINK_TARGET).replace("/", "\\")
    old = cell.value if isinstance(cell.value, str) else ""
    if LINK_TEXT:
        text = render(LINK_TEXT)
    elif old.upper().startswith("=HYPERLINK("):
        m = HYPERLINK_FORMULA_TEXT.search(old)  # =HYPERLINK("路徑","顯示文字") 取顯示文字
        text = update_week_tokens(m.group(1), week) if m else target.rstrip("\\").split("\\")[-1]
    elif old.strip():
        text = update_week_tokens(old, week)
    else:
        text = target.rstrip("\\").split("\\")[-1]  # 空白格顯示最後一層資料夾名稱，例: BOS
    cell.value = text
    cell.hyperlink = target
    f = cell.font
    if not f.underline:  # 沒有連結樣式時改成藍字底線，其他字型設定保留
        cell.font = Font(name=f.name, size=f.size, bold=f.bold, italic=f.italic,
                         underline="single", color="FF0563C1")


def finalize_weekly_sheet(ws, cols, host_rows, filled_rows, week):
    """範本模式：沒有本週 CSV 的伺服器填 O O、防毒填 O、連結改成本週；回傳沒有 CSV 而填 O O 的伺服器"""
    no_csv = []
    for name, r in sorted(host_rows.items(), key=lambda x: x[1]):
        host = str(ws.cell(row=r, column=1).value).strip()
        if not HOSTNAME_RE.match(host):
            continue  # 分隔列等非伺服器名稱
        if r not in filled_rows:
            wrote = set_cell(ws, r, cols[WARN_HEADER], "O")
            wrote = set_cell(ws, r, cols[ERROR_HEADER], "O") or wrote
            if wrote:
                set_cell(ws, r, cols[NOTE_HEADER], None)
                no_csv.append(host)
        if ANTIVIRUS_HEADER:
            set_cell(ws, r, cols[ANTIVIRUS_HEADER], "O")
        if LINK_HEADER:
            cell = ws.cell(row=r, column=cols[LINK_HEADER])
            if is_shaded(cell):
                continue
            if LINK_TARGET:
                set_link(cell, host, week)
                continue
            if isinstance(cell.value, str):
                cell.value = update_week_tokens(cell.value, week)
            if cell.hyperlink is not None:
                if cell.hyperlink.target:
                    cell.hyperlink.target = update_week_tokens(cell.hyperlink.target, week)
                if cell.hyperlink.display:
                    cell.hyperlink.display = update_week_tokens(cell.hyperlink.display, week)
    return no_csv


# ---------- 問題追蹤表 ----------

def parse_filename(path):
    """檔名拆出 (伺服器名, 日期)，例: 20261015_P3-WFA-AP01V.csv → (P3-WFA-AP01V, 2026-10-15)"""
    name = os.path.splitext(os.path.basename(path))[0].strip()
    m = re.match(r"^(\d{8})_(.+)$", name)
    if m:
        try:
            return m.group(2).strip(), datetime.strptime(m.group(1), "%Y%m%d").date()
        except ValueError:
            pass
    return re.sub(r"^\d+_", "", name), date.today()


def week_label(d):
    """ISO 週次，例: 2026-10-15 → W2642"""
    year, week, _ = d.isocalendar()
    return f"W{year % 100:02d}{week:02d}"


def get_track_sheet(wb):
    if not TRACK_SHEET:
        return wb.worksheets[0]
    if TRACK_SHEET not in wb.sheetnames:
        raise KeyError(f"Tracking file: sheet '{TRACK_SHEET}' not found, sheets: {wb.sheetnames}")
    return wb[TRACK_SHEET]


def track_columns(ws):
    """依標題找出追蹤表各欄欄號：伺服器、週別、CSV 對應欄、可忽略欄（可省略）"""
    headers = {re.sub(r"\s+", "", str(c.value)): c.column
               for c in ws[TRACK_HEADER_ROW] if c.value is not None}
    need = [HOST_HEADER, WEEK_HEADER] + list(TRACK_COLUMNS)
    missing = [h for h in need if re.sub(r"\s+", "", h) not in headers]
    if missing:
        raise KeyError(f"Tracking sheet: headers {missing} not found, headers: {list(headers)}")
    cols = {h: headers[re.sub(r"\s+", "", h)] for h in need}
    ignore_col = headers.get(re.sub(r"\s+", "", IGNORE_HEADER)) if IGNORE_HEADER else None
    return cols, ignore_col


def delete_existing(ws, cols, host, week):
    """刪除同一伺服器、同一週別的舊資料（由下往上刪才不會錯位）"""
    removed = 0
    for r in range(ws.max_row, TRACK_HEADER_ROW, -1):
        h = ws.cell(row=r, column=cols[HOST_HEADER]).value
        w = ws.cell(row=r, column=cols[WEEK_HEADER]).value
        if h is not None and str(h).strip().upper() == host.upper() and str(w).strip() == week:
            ws.delete_rows(r)
            removed += 1
    return removed


def last_data_row(ws, cols):
    """伺服器欄最後一個有值的列號（可忽略欄預先拉好的公式不算資料）"""
    for r in range(ws.max_row, TRACK_HEADER_ROW, -1):
        if ws.cell(row=r, column=cols[HOST_HEADER]).value not in (None, ""):
            return r
    return TRACK_HEADER_ROW


def cell_value(v):
    """CSV 值轉成儲存格值；純數字（例: 事件 Id）存成數字，避免 Excel 顯示「以文字儲存的數字」"""
    if pd.isna(v):
        return None
    s = str(v)
    return int(s) if s.strip().isdigit() and len(s.strip()) < 15 else s


def paste_csv(ws, cols, df, host, week):
    removed = delete_existing(ws, cols, host, week)
    r = last_data_row(ws, cols) + 1
    for _, row in df.iterrows():
        ws.cell(row=r, column=cols[HOST_HEADER]).value = host
        ws.cell(row=r, column=cols[WEEK_HEADER]).value = week
        for header, csv_col in TRACK_COLUMNS.items():
            ws.cell(row=r, column=cols[header]).value = cell_value(row[csv_col]) if csv_col in df.columns else None
        r += 1
    print(f"  Tracking sheet: {week} pasted {len(df)} rows" + (f" (replaced {removed} old rows)" if removed else ""))


def fill_ignore_marks(ws, cols, ignore_col, ignore_ids):
    """可忽略欄：每筆資料填公式（追蹤表含可忽略清單時）或直接填 V；刪列後也一併校正"""
    if not ignore_col:
        return
    id_col = cols[TRACK_ID_HEADER]
    id_letter = openpyxl.utils.get_column_letter(id_col)
    src = None
    wb = ws.parent
    if IGNORE_SHEET in wb.sheetnames:
        ign = wb[IGNORE_SHEET]
        for c in ign[1]:
            if c.value is not None and str(c.value).strip() == IGNORE_COLUMN:
                L = openpyxl.utils.get_column_letter(c.column)
                src = f"'{IGNORE_SHEET}'!${L}$2:${L}$25555"
    for r in range(TRACK_HEADER_ROW + 1, last_data_row(ws, cols) + 1):
        cell = ws.cell(row=r, column=ignore_col)
        if src:
            cell.value = f'=IF(COUNTIF({src},"="&{id_letter}{r})>0,"V","")'
        else:
            cell.value = "V" if norm(ws.cell(row=r, column=id_col).value) in ignore_ids else None


# ---------- 主程式 ----------

def account():
    user = os.environ.get("USERNAME") or os.environ.get("USER") or "?"
    domain = os.environ.get("USERDOMAIN")
    return f"{domain}\\{user}" if domain else user


def save(wb, path):
    try:
        folder = os.path.dirname(path)
        if folder and not os.path.isdir(folder):
            os.makedirs(folder, exist_ok=True)
        wb.save(path)
    except PermissionError as e:
        print(f"[ERROR] Cannot write {path}: {e}")
        print(f"        Close the file if it is open in Excel. Otherwise account {account()} needs Modify permission on the folder.")
        sys.exit(1)
    except OSError as e:
        print(f"[ERROR] Cannot write {path}: {e}")
        print(f"        Check the path and that account {account()} can access and write to the folder.")
        sys.exit(1)
    print(f"Saved: {path}")


def parse_args(argv):
    """[--week W2642] [csv資料夾] [檢核表xlsx] [追蹤表xlsx]"""
    args, week = list(argv), None
    if "--week" in args:
        k = args.index("--week")
        week = args[k + 1].upper() if k + 1 < len(args) else ""
        del args[k:k + 2]
        if not re.fullmatch(r"W\d{4}", week):
            print(f"[ERROR] --week must look like W2642, got '{week}'")
            sys.exit(1)
    return args, week


def main():
    args, week = parse_args(sys.argv[1:])
    csv_dir = args[0] if len(args) > 0 else CSV_DIR
    check_xlsx = args[1] if len(args) > 1 else CHECK_XLSX
    track_xlsx = args[2] if len(args) > 2 else TRACK_XLSX
    template_mode = bool(TEMPLATE_XLSX)
    # 目標週：預設今天所在的 ISO 週；只處理這一週的 CSV，不會誤用上週的
    week = week or week_label(date.today())

    if not os.path.isdir(csv_dir):
        print(f"[ERROR] CSV folder not found: {csv_dir} (check paths.csv_dir in config.json)")
        sys.exit(1)
    all_csv = sorted(glob.glob(os.path.join(csv_dir, "*.csv")))
    latest = {}  # 伺服器 → (日期, 路徑)；同一週有多次時取最新日期
    for p in all_csv:
        host, d = parse_filename(p)
        if week_label(d) == week and (host.upper() not in latest or d > latest[host.upper()][0]):
            latest[host.upper()] = (d, p)
    csv_files = sorted(p for _, p in latest.values())
    other = len(all_csv) - len(csv_files)
    print(f"Week {week}: {len(csv_files)} CSV files" + (f", skipped {other} files of other weeks/older runs" if other else ""))
    if not csv_files and not template_mode:
        print(f"No CSV files for {week} in {csv_dir}")
        return

    track_xlsx = week_path(track_xlsx, week)
    track_output = week_path(TRACK_OUTPUT or track_xlsx, week)
    if template_mode:
        # 當週檔已存在（重跑）就以它為底，保留手動填的內容；否則從範本產生
        check_output = week_path(OUTPUT_XLSX or check_xlsx, week)
        check_src = check_output if os.path.isfile(check_output) else week_path(TEMPLATE_XLSX, week)
    else:
        check_src = week_path(check_xlsx, week)
        check_output = week_path(OUTPUT_XLSX or check_src, week)
    for label, key, p in (("check sheet", "check_sheet.template" if template_mode else "paths.check_xlsx", check_src),
                          ("tracking file", "paths.track_xlsx", track_xlsx)):
        if not os.path.isfile(p):
            print(f"[ERROR] {label} not found: {p} (check {key} in config.json)")
            sys.exit(1)
    print(f"CSV folder   : {csv_dir}")
    if template_mode:
        print(f"Check sheet  : {check_src}" + ("  (existing file of this week)" if check_src == check_output else "  (template)"))
        print(f"Output       : {check_output}")
    else:
        print(f"Check sheet  : {check_src}")
    print(f"Tracking file: {track_xlsx}")

    ignore_ids = load_ignore_ids(track_xlsx)
    print(f"Ignorable event IDs: {len(ignore_ids)}")

    check_wb = openpyxl.load_workbook(check_src)
    check_ws = check_wb[CHECK_SHEET] if CHECK_SHEET else check_wb.worksheets[0]
    cols = find_columns(check_ws, (ANTIVIRUS_HEADER, LINK_HEADER) if template_mode else ())
    host_rows = find_host_rows(check_ws)
    filled_rows = set()

    track_wb = openpyxl.load_workbook(track_xlsx)
    track_ws = get_track_sheet(track_wb)
    track_cols, ignore_col = track_columns(track_ws)

    problems = []  # (檔名, 狀況) 最後彙總
    for path in csv_files:
        host, d = parse_filename(path)
        fname = os.path.basename(path)
        print(host)
        try:
            df, skipped, unknown = keep_warn_error(read_csv_any(path))
        except Exception as e:
            print(f"  [ERROR] {e}")
            problems.append((fname, f"ERROR: cannot read CSV, skipped ({e})"))
            continue
        print(f"  Errors/warnings: {len(df)}" + (f", skipped {skipped} information/other rows" if skipped else ""))
        if unknown:
            msg = (f"unknown level values {unknown} were skipped - the CSV may be garbled "
                   f"(re-collect with the latest EventLogCheck.ps1)")
            print(f"  [WARN] {msg}")
            problems.append((fname, f"WARN: {msg}"))

        r = host_rows.get(host.upper())
        if r is None:
            print("  [SKIP] Server not found in first column of check sheet")
            problems.append((fname, "SKIP: server not in check sheet, O/X not filled (tracking sheet still pasted)"))
        else:
            try:
                fill_check_row(check_ws, cols, r, df, ignore_ids, host)
                filled_rows.add(r)
            except Exception as e:
                print(f"  [ERROR] Check sheet: {e}")
                problems.append((fname, f"ERROR: O/X not filled ({e})"))

        paste_csv(track_ws, track_cols, df, host, week_label(d))

    if TRACK_ID_HEADER:
        fill_ignore_marks(track_ws, track_cols, ignore_col, ignore_ids)

    if template_mode:
        no_csv = finalize_weekly_sheet(check_ws, cols, host_rows, filled_rows, week)
        print(f"Check sheet: antivirus = O, links updated to {week}")
        if no_csv:
            print(f"Servers without a {week} CSV, event columns set to O O: {', '.join(no_csv)}")

    save(check_wb, check_output)
    save(track_wb, track_output)

    print(f"\n===== Done: {len(csv_files)} CSV files, {len(problems)} with problems =====")
    for fname, msg in problems:
        print(f"  {fname}: {msg}")


class _Tee:
    """畫面輸出同時寫進 log 檔"""
    def __init__(self, *streams):
        self.streams = streams

    def write(self, s):
        for st in self.streams:
            try:
                st.write(s)
            except (UnicodeEncodeError, OSError, ValueError):
                # 畫面顯示不了的字（或視窗已關）不影響 log
                pass

    def flush(self):
        for st in self.streams:
            try:
                st.flush()
            except (OSError, ValueError):
                pass


def run():
    """執行 main，畫面輸出全部寫進 log；未預期的錯誤只在畫面顯示一行，完整錯誤內容記在 log"""
    os.makedirs(LOG_DIR, exist_ok=True)
    log_path = os.path.join(LOG_DIR, f"{datetime.now():%Y%m%d}.log")
    console = sys.stdout
    code = 0
    # 同一天執行多次都附加在同一個 log 檔
    # 新檔才寫 BOM，舊版記事本開啟中文才不會亂碼
    enc = "utf-8" if os.path.exists(log_path) else "utf-8-sig"
    with open(log_path, "a", encoding=enc) as log:
        sys.stdout = _Tee(console, log)
        try:
            print(f"\n===== fill_check start {datetime.now():%Y-%m-%d %H:%M:%S} =====")
            main()
        except SystemExit as e:
            code = e.code or 0
        except Exception as e:
            code = 1
            print(f"[ERROR] Aborted, Excel files NOT saved: {type(e).__name__}: {e}")
            log.write("\n" + traceback.format_exc())
        finally:
            print(f"End {datetime.now():%Y-%m-%d %H:%M:%S}, log: {log_path}")
            sys.stdout = console
    sys.exit(code)


if __name__ == "__main__":
    run()
