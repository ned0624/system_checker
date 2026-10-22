# -*- coding: utf-8 -*-
"""CI: fill_check template mode end to end.

Builds a template check sheet + tracking file + CSVs in a temp folder, runs fill_check.py --week W2642
twice (second run must keep manual edits) and checks the generated weekly file.
Usage: python ci_template_test.py <empty temp folder>
"""
import json
import os
import shutil
import subprocess
import sys

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color
from openpyxl.worksheet.hyperlink import Hyperlink

tmp = os.path.abspath(sys.argv[1])
os.makedirs(tmp, exist_ok=True)
app = os.path.join(tmp, "app")
os.makedirs(app, exist_ok=True)
shutil.copy("fill_check.py", app)

GRAY = PatternFill("solid", fgColor="FF808080")
BLACK = PatternFill("solid", fgColor="FF000000")
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
THEME_GRAY = PatternFill("solid", fgColor=Color(theme=0, tint=-0.25))
WHITE = PatternFill("solid", fgColor="FFFFFFFF")
SHARE = r"\\xintec.com.tw\tsm\Supporting_Org\Dept\System Check"

# ---- template ----
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Server"
ws.append(["Host Name", "Checker", "事件檢視器\n系統錯誤", "事件檢視器\n警告事件", "防毒軟體", "CPU", "連結", "備註"])
for host in ("SRV-A", "SRV-B", "SRV-C", "SRV-D"):
    ws.append([host, "Ned"])
ws.append(["(季檢,下次 10/1, 2027/1/7)"])   # row 6: yellow merged separator
ws.append(["SRV-E", "Ned"])                  # row 7
for c in "CDE":
    ws[f"{c}4"].fill = GRAY                  # SRV-C: event + antivirus gray
ws["E5"].fill = THEME_GRAY                   # SRV-D: antivirus theme gray
ws["C7"].fill = BLACK                        # SRV-E: error column black
ws["D2"].fill = WHITE                        # explicit white is still fillable
for c in "ABCDEFGH":
    ws[f"{c}6"].fill = YELLOW
ws.merge_cells("A6:H6")
ws["G2"] = "SRV-A_Evidence_W40"
ws["G2"].hyperlink = Hyperlink(ref="G2", target=f"file:///{SHARE}\\W2640\\Reference\\BOS - SRV-A_Evidence_W40.docx")
ws["G3"] = f'=HYPERLINK("{SHARE}\\W2640\\Reference\\BOS - SRV-B_Evidence_W40.docx","SRV-B_Evidence_W40")'
tpl = os.path.join(tmp, "template.xlsx")
wb.save(tpl)

# ---- tracking file ----
wb = openpyxl.Workbook()
ig = wb.active
ig.title = "可忽略事件ID"
ig.append(["可忽略事件ID"])
ig.append([10016])
tr = wb.create_sheet("事件檢視器_BOS")
tr.append(["Server Name", "周", "錯誤/警告", "事件識別碼", "記錄檔名稱", "事件內容", "可能造成的影響", "Action", "可忽略"])
track = os.path.join(tmp, "track.xlsx")
wb.save(track)

# ---- CSVs: SRV-A this week (W2642 = 2026-10-12..18), SRV-B only last week ----
csv_dir = os.path.join(tmp, "csv")
os.makedirs(csv_dir, exist_ok=True)
hdr = "LevelDisplayName,Id,LogName,Message,TimeCreated,Count\n"
open(os.path.join(csv_dir, "20261015_SRV-A.csv"), "w", encoding="utf-8-sig").write(
    hdr + "錯誤,1801,System,boot cert,2026-10-15 10:00:00,1\n警告,10016,System,DCOM,2026-10-15 10:00:00,3\n")
open(os.path.join(csv_dir, "20261008_SRV-B.csv"), "w", encoding="utf-8-sig").write(
    hdr + "錯誤,999,System,last week error,2026-10-08 10:00:00,1\n")

out_pattern = os.path.join(tmp, "out", "{week}", "BOS_{week}.xlsx")
cfg = {
    "paths": {"csv_dir": csv_dir, "check_xlsx": tpl, "track_xlsx": track},
    "ignore_list": {"sheet": "可忽略事件ID", "id_column": "可忽略事件ID"},
    "csv": {"id_column": "Id", "level_column": "LevelDisplayName", "desc_column": "Message"},
    "check_sheet": {"sheet": None, "header_row": 1, "warn_header": "事件檢視器\n警告事件",
                    "error_header": "事件檢視器\n系統錯誤", "note_header": "備註", "output": out_pattern,
                    "template": tpl, "antivirus_header": "防毒軟體", "link_header": "連結"},
    "track": {"sheet": "事件檢視器_BOS", "header_row": 1, "host_header": "Server Name", "week_header": "周",
              "ignore_header": "可忽略", "output": None},
    "log_dir": os.path.join(tmp, "logs"),
}
json.dump(cfg, open(os.path.join(app, "config.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=2)


def run():
    r = subprocess.run([sys.executable, os.path.join(app, "fill_check.py"), "--week", "W2642"],
                       capture_output=True, text=True, encoding="utf-8", errors="replace",
                       env=dict(os.environ, PYTHONIOENCODING="utf-8"))
    print(r.stdout, r.stderr)
    assert r.returncode == 0, r.returncode


run()
out = os.path.join(tmp, "out", "W2642", "BOS_W2642.xlsx")
assert os.path.isfile(out), "weekly file not created in W2642 folder"
ws = openpyxl.load_workbook(out).active
v = lambda ref: ws[ref].value

# SRV-A (row 2): error X (1801), warning O (10016 ignorable), note has 1801 only
assert (v("C2"), v("D2")) == ("X", "O"), (v("C2"), v("D2"))
assert "1801" in v("H2") and "10016" not in v("H2"), v("H2")
# SRV-B (row 3): only last week's CSV -> treated as no CSV -> O O
assert (v("C3"), v("D3"), v("H3")) == ("O", "O", None), (v("C3"), v("D3"), v("H3"))
# SRV-C (row 4): gray cells untouched
assert (v("C4"), v("D4"), v("E4")) == (None, None, None), (v("C4"), v("D4"), v("E4"))
# SRV-D (row 5): theme-gray antivirus untouched, events O O
assert v("E5") is None and (v("C5"), v("D5")) == ("O", "O"), (v("E5"), v("C5"), v("D5"))
# separator row 6 untouched
assert all(v(f"{c}6") is None for c in "BCDEFGH"), [v(f"{c}6") for c in "BCDEFGH"]
# SRV-E (row 7): black error cell untouched, warning O, antivirus O
assert v("C7") is None and v("D7") == "O" and v("E7") == "O", (v("C7"), v("D7"), v("E7"))
# antivirus O on normal rows
assert [v(f"E{r}") for r in (2, 3)] == ["O", "O"]
# links updated
assert v("G2") == "SRV-A_Evidence_W42", v("G2")
assert "\\W2642\\" in ws["G2"].hyperlink.target and "_W42.docx" in ws["G2"].hyperlink.target, ws["G2"].hyperlink.target
assert "\\W2642\\" in v("G3") and "SRV-B_Evidence_W42" in v("G3") and "W2640" not in v("G3"), v("G3")
# tracking sheet: only this week's CSV pasted
trs = openpyxl.load_workbook(track)["事件檢視器_BOS"]
pasted = [(r[0], r[1]) for r in trs.iter_rows(min_row=2, values_only=True) if r[0]]
assert pasted and all(p == ("SRV-A", "W2642") for p in pasted), pasted

# second run keeps manual edits made in the weekly file
wb = openpyxl.load_workbook(out)
wb.active["F2"] = "manual 30%"
wb.save(out)
run()
ws = openpyxl.load_workbook(out).active
assert ws["F2"].value == "manual 30%", ws["F2"].value
assert ws["G2"].value == "SRV-A_Evidence_W42" and ws["C2"].value == "X"
print("TEMPLATE TEST OK")
